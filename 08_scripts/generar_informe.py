#!/usr/bin/env python3
"""Genera HTML y DOCX desde el informe Markdown y conserva la plantilla académica."""
from __future__ import annotations

import html
import re
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

RAIZ = Path(__file__).resolve().parents[1]
INFORME = RAIZ / "07_informe"
MARKDOWN = INFORME / "informe_laboratorio_07.md"
PLANTILLA = Path("/home/carlos-daniel/Descargas/STI Laboratorio Hardering - Ayala.docx")
LOGO = INFORME / "recursos" / "logo_institucional.png"
SALIDA_DOCX = INFORME / "informe_laboratorio_07.docx"
SALIDA_HTML = INFORME / "informe_laboratorio_07.html"
ANCHO_CONTENIDO = 6.3


def establecer_fuente(run, nombre="Times New Roman", tamano=12, negrita=None, cursiva=None):
    run.font.name = nombre
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), nombre)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), nombre)
    run.font.size = Pt(tamano)
    if negrita is not None:
        run.bold = negrita
    if cursiva is not None:
        run.italic = cursiva


def sombrear(celda, color="E7E6E6"):
    tc_pr = celda._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def margen_celda(celda, top=65, start=75, bottom=65, end=75):
    tc = celda._tc
    tc_pr = tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for lado, valor in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        nodo = mar.find(qn(f"w:{lado}"))
        if nodo is None:
            nodo = OxmlElement(f"w:{lado}")
            mar.append(nodo)
        nodo.set(qn("w:w"), str(valor))
        nodo.set(qn("w:type"), "dxa")


def agregar_break_linea(parrafo):
    parrafo.add_run().add_break()


def agregar_texto_enriquecido(parrafo, texto, tamano=12, color=None):
    """Inserta negritas, cursivas y saltos <br> con formato Times New Roman."""
    texto = texto.replace("<br>", "\n").replace("<br/>", "\n")
    tokens = re.split(r"(`[^`]+`|\*\*.*?\*\*|\*.*?\*|\n)", texto)
    for token in tokens:
        if not token:
            continue
        if token == "\n":
            agregar_break_linea(parrafo)
            continue
        if token.startswith("`") and token.endswith("`"):
            run = parrafo.add_run(token[1:-1])
            establecer_fuente(run, nombre="Courier New", tamano=tamano)
        elif token.startswith("**") and token.endswith("**"):
            run = parrafo.add_run(token[2:-2])
            establecer_fuente(run, tamano=tamano, negrita=True)
        elif token.startswith("*") and token.endswith("*"):
            run = parrafo.add_run(token[1:-1])
            establecer_fuente(run, tamano=tamano, cursiva=True)
        else:
            run = parrafo.add_run(token)
            establecer_fuente(run, tamano=tamano)
        if color:
            run.font.color.rgb = RGBColor.from_string(color)


def preparar_estilos(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    for nombre, tamano in (("Heading 1", 14), ("Heading 2", 12)):
        estilo = doc.styles[nombre]
        estilo.font.name = "Times New Roman"
        estilo._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        estilo._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        estilo.font.size = Pt(tamano)
        estilo.font.bold = True
        estilo.font.color.rgb = RGBColor(0, 0, 0)
        estilo.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        estilo.paragraph_format.space_before = Pt(10 if nombre == "Heading 1" else 6)
        estilo.paragraph_format.space_after = Pt(4)
        estilo.paragraph_format.keep_with_next = True


def limpiar_cuerpo(doc):
    cuerpo = doc._element.body
    sectpr = cuerpo.sectPr
    for elemento in list(cuerpo):
        if elemento is not sectpr:
            cuerpo.remove(elemento)
    for rid, relacion in list(doc.part.rels.items()):
        if relacion.reltype == RT.IMAGE:
            doc.part.drop_rel(rid)


def normalizar_twips_seccion(doc):
    # La plantilla almacena los márgenes equivalentes a 2,5 cm como decimales;
    # Word admite el redondeo al twip entero más próximo (diferencia < 0,01 mm).
    for seccion in doc.sections:
        pgmar = seccion._sectPr.pgMar
        for atributo in ("top", "bottom", "left", "right", "header", "footer", "gutter"):
            clave = qn(f"w:{atributo}")
            valor = pgmar.get(clave)
            if valor is not None and "." in valor:
                pgmar.set(clave, str(round(float(valor))))


def agregar_portada(doc):
    seccion = doc.sections[0]
    seccion.different_first_page_header_footer = True

    def linea(texto, tamano, after=1):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.keep_together = True
        run = p.add_run(texto)
        establecer_fuente(run, tamano=tamano, negrita=True)
        return p

    linea("UNIVERSIDAD PRIVADA DE TACNA", 18, 3)
    linea("FACULTAD DE INGENIERÍA", 16, 2)
    linea("ESCUELA DE INGENIERÍA DE SISTEMAS", 14, 1)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(10)
    p.add_run().add_picture(str(LOGO), width=Inches(1.35))
    linea("Guía Práctica de Laboratorio N.° 07", 25, 2)
    linea("“Análisis de Malware”", 19, 12)
    linea("Que se presenta para el curso:", 12, 2)
    linea("Seguridad de Tecnologías de Información", 14, 10)
    linea("Estudiante: Carlos Daniel Ayala Ramos", 13, 4)
    linea("Docente: Dr. Renzo Alberto Taco Coayla", 13, 12)
    linea("TACNA – PERÚ", 12, 2)
    linea("2026", 12, 0)
    doc.add_page_break()


def establecer_repeticion_fila(fila):
    tr_pr = fila._tr.get_or_add_trPr()
    encabezado = OxmlElement("w:tblHeader")
    encabezado.set(qn("w:val"), "true")
    tr_pr.append(encabezado)


def establecer_no_dividir_fila(fila):
    tr_pr = fila._tr.get_or_add_trPr()
    nodo = OxmlElement("w:cantSplit")
    tr_pr.append(nodo)


def establecer_bordes_tabla(tabla):
    tbl_pr = tabla._tbl.tblPr
    bordes = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borde = OxmlElement(f"w:{lado}")
        borde.set(qn("w:val"), "single")
        borde.set(qn("w:sz"), "4")
        borde.set(qn("w:space"), "0")
        borde.set(qn("w:color"), "666666")
        bordes.append(borde)
    tbl_pr.append(bordes)


def parsear_fila_tabla(linea):
    return [celda.strip() for celda in linea.strip().strip("|").split("|")]


def tamano_tabla(filas, ancho):
    if not filas:
        return 8.2
    encabezado = " ".join(filas[0]).lower()
    if "sha-256" in encabezado and "ssdeep" in encabezado:
        return 6.6
    if "recurso" in encabezado or "estado" in encabezado:
        return 8.5 if len(filas[0]) > 3 else 10.0
    return 8.0 if len(filas[0]) >= 5 else 9.0


def agregar_tabla(doc, filas):
    if len(filas) < 2:
        return
    columnas = max(len(f) for f in filas)
    tabla = doc.add_table(rows=1, cols=columnas)
    establecer_bordes_tabla(tabla)
    tabla.autofit = False
    tabla.alignment = 1
    ancho_fuente = tamano_tabla(filas, ANCHO_CONTENIDO)
    pesos = {
        2: [1.75, 4.55],
        4: [1.15, 1.05, 1.95, 2.15],
        5: [0.55, 2.1, 0.65, 1.3, 1.7],
        6: [0.42, 1.85, 0.55, 1.2, 1.25, 1.03],
    }.get(columnas, [ANCHO_CONTENIDO / columnas] * columnas)
    total = sum(pesos)
    pesos = [ANCHO_CONTENIDO * p / total for p in pesos]
    for i, ancho in enumerate(pesos):
        tabla.columns[i].width = Inches(ancho)

    for indice, datos in enumerate(filas):
        fila = tabla.rows[0] if indice == 0 else tabla.add_row()
        establecer_no_dividir_fila(fila)
        if indice == 0:
            establecer_repeticion_fila(fila)
        for col in range(columnas):
            celda = fila.cells[col]
            celda.width = Inches(pesos[col])
            celda.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margen_celda(celda)
            celda.text = ""
            p = celda.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            contenido = datos[col] if col < len(datos) else ""
            if columnas == 4 and indice > 0 and col == 3 and len(contenido) > 36 and "<br>" not in contenido:
                contenido = contenido[:26] + "<br>" + contenido[26:52] + "<br>" + contenido[52:]
            agregar_texto_enriquecido(p, contenido, tamano=ancho_fuente)
            for run in p.runs:
                if indice == 0:
                    run.bold = True
                if columnas == 4 and indice > 0 and col in (1, 2, 3):
                    run.font.name = "Courier New"
                    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Courier New")
                    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Courier New")
                    run.font.size = Pt(ancho_fuente)
            if indice == 0:
                sombrear(celda)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def agregar_imagen(doc, ruta_relativa):
    ruta = (INFORME / ruta_relativa).resolve()
    if not ruta.exists():
        raise FileNotFoundError(f"No se encontró la evidencia: {ruta}")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_together = True
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(ruta), width=Inches(2.4))


def agregar_parrafo(doc, lineas, referencias=False, izquierda=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if referencias or izquierda else WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.widow_control = True
    if referencias:
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.28)
    for n, linea in enumerate(lineas):
        hard_break = linea.endswith("  ")
        if n:
            p.add_run(" ")
        agregar_texto_enriquecido(p, linea.rstrip(), tamano=12)
        if hard_break:
            agregar_break_linea(p)
    return p


def convertir_docx(texto):
    if not PLANTILLA.exists():
        raise FileNotFoundError(f"No se encontró la plantilla: {PLANTILLA}")
    if not LOGO.exists():
        raise FileNotFoundError(f"No se encontró el logotipo extraído: {LOGO}")
    doc = Document(str(PLANTILLA))
    normalizar_twips_seccion(doc)
    limpiar_cuerpo(doc)
    preparar_estilos(doc)
    doc.core_properties.title = "Guía Práctica de Laboratorio N.° 07: Análisis de Malware"
    doc.core_properties.subject = "Seguridad de Tecnologías de Información"
    doc.core_properties.author = "Carlos Daniel Ayala Ramos"
    agregar_portada(doc)

    lineas = texto.splitlines()
    inicio = next((i + 1 for i, x in enumerate(lineas) if x.strip() == "[[PAGEBREAK]]"), 0)
    lineas = lineas[inicio:]
    i = 0
    en_referencias = False
    while i < len(lineas):
        actual = lineas[i].rstrip()
        if not actual.strip():
            i += 1
            continue
        if actual.strip() == "[[PAGEBREAK]]":
            doc.add_page_break()
            i += 1
            continue
        imagen = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", actual.strip())
        if imagen:
            agregar_imagen(doc, imagen.group(2))
            i += 1
            continue
        if actual.lstrip().startswith("|"):
            filas = []
            while i < len(lineas) and lineas[i].lstrip().startswith("|"):
                fila = parsear_fila_tabla(lineas[i])
                if not all(re.fullmatch(r":?-{3,}:?", celda or "") for celda in fila):
                    filas.append(fila)
                i += 1
            agregar_tabla(doc, filas)
            continue
        if actual.startswith("#"):
            nivel = len(actual) - len(actual.lstrip("#"))
            titulo = actual[nivel:].strip()
            if nivel == 1:
                en_referencias = titulo.startswith("6. Referencias bibliográficas")
            p = doc.add_paragraph(style="Heading 1" if nivel == 1 else "Heading 2")
            agregar_texto_enriquecido(p, titulo, tamano=14 if nivel == 1 else 12)
            i += 1
            continue
        if actual.startswith("- "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.left_indent = Inches(0.28)
            p.paragraph_format.first_line_indent = Inches(-0.15)
            viñeta = p.add_run("• ")
            establecer_fuente(viñeta, tamano=11.5)
            agregar_texto_enriquecido(p, actual[2:].strip(), tamano=11.5)
            i += 1
            continue
        if re.match(r"^\d+\.\s", actual):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            agregar_texto_enriquecido(p, actual, tamano=11.5)
            i += 1
            continue

        grupo = [actual]
        i += 1
        while i < len(lineas):
            siguiente = lineas[i].rstrip()
            if not siguiente.strip() or siguiente.strip() == "[[PAGEBREAK]]" or siguiente.startswith("#") or siguiente.lstrip().startswith("|") or siguiente.startswith("- ") or re.match(r"^\d+\.\s", siguiente) or re.match(r"!\[([^\]]*)\]\(([^)]+)\)", siguiente.strip()):
                break
            grupo.append(siguiente)
            i += 1
        contenido_plano = " ".join(x.rstrip() for x in grupo).strip()
        alinear_izquierda = contenido_plano.startswith((
            "Se revisaron dos informes públicos",
            "El segundo reporte, titulado",
        ))
        parrafo = agregar_parrafo(doc, grupo, referencias=en_referencias, izquierda=alinear_izquierda)
        if contenido_plano.startswith("**Figura "):
            parrafo.paragraph_format.keep_together = True

    doc.save(str(SALIDA_DOCX))


def inline_html(texto):
    texto = html.escape(texto, quote=False)
    texto = re.sub(r"`([^`]+)`", r"<code>\1</code>", texto)
    texto = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", texto)
    texto = re.sub(r"\*(.+?)\*", r"<em>\1</em>", texto)
    texto = re.sub(r"(https?://[^\s<]+)", r'<a href="\1">\1</a>', texto)
    return texto.replace("<br>", "<br />")


def convertir_html(texto):
    lineas = texto.splitlines()
    cuerpo = []
    i = 0
    while i < len(lineas):
        actual = lineas[i].rstrip()
        if not actual.strip():
            i += 1
            continue
        if actual.strip() == "[[PAGEBREAK]]":
            cuerpo.append('<div class="pagebreak"></div>')
            i += 1
            continue
        imagen = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", actual.strip())
        if imagen:
            cuerpo.append(f'<figure><img src="{html.escape(imagen.group(2), quote=True)}" alt="{html.escape(imagen.group(1), quote=True)}" /></figure>')
            i += 1
            continue
        if actual.lstrip().startswith("|"):
            filas = []
            while i < len(lineas) and lineas[i].lstrip().startswith("|"):
                fila = parsear_fila_tabla(lineas[i])
                if not all(re.fullmatch(r":?-{3,}:?", celda or "") for celda in fila):
                    filas.append(fila)
                i += 1
            cuerpo.append("<table>")
            for n, fila in enumerate(filas):
                tag = "th" if n == 0 else "td"
                cuerpo.append("<tr>" + "".join(f"<{tag}>{inline_html(celda)}</{tag}>" for celda in fila) + "</tr>")
            cuerpo.append("</table>")
            continue
        if actual.startswith("#"):
            nivel = min(len(actual) - len(actual.lstrip("#")), 6)
            cuerpo.append(f"<h{nivel}>{inline_html(actual[nivel:].strip())}</h{nivel}>")
            i += 1
            continue
        if actual.startswith("- "):
            cuerpo.append("<ul>")
            while i < len(lineas) and lineas[i].startswith("- "):
                cuerpo.append(f"<li>{inline_html(lineas[i][2:].strip())}</li>")
                i += 1
            cuerpo.append("</ul>")
            continue
        grupo = [actual]
        i += 1
        while i < len(lineas):
            siguiente = lineas[i].rstrip()
            if not siguiente.strip() or siguiente.strip() == "[[PAGEBREAK]]" or siguiente.startswith("#") or siguiente.lstrip().startswith("|") or siguiente.startswith("- ") or re.match(r"!\[([^\]]*)\]\(([^)]+)\)", siguiente.strip()):
                break
            grupo.append(siguiente)
            i += 1
        cuerpo.append("<p>" + inline_html(" ".join(x.rstrip() for x in grupo)) + "</p>")

    cuerpo_html = "\n".join(cuerpo)
    pagina = f"""<!doctype html>
<html lang="es">
<head>
<meta charset="utf-8" />
<meta name="viewport" content="width=device-width, initial-scale=1" />
<title>Laboratorio N.° 07 — Análisis de Malware</title>
<style>
@page {{ size: A4; margin: 2.5cm; }}
body {{ max-width: 17cm; margin: 2.5cm auto; color: #111; font-family: 'Times New Roman', serif; font-size: 12pt; line-height: 1.2; }}
h1 {{ font-size: 14pt; margin: 1.1em 0 .35em; }}
h2 {{ font-size: 12pt; margin: .8em 0 .3em; }}
h1:first-of-type {{ font-size: 20pt; }}
p {{ text-align: justify; margin: 0 0 .7em; }}
table {{ border-collapse: collapse; width: 100%; font-size: 9pt; margin: .6em 0 1em; }}
th, td {{ border: 1px solid #555; padding: .35em; vertical-align: top; }}
th {{ background: #e7e6e6; }}
figure {{ text-align: center; margin: 1em 0; }}
figure img {{ max-height: 17cm; max-width: 2.65in; object-fit: contain; }}
.pagebreak {{ page-break-after: always; break-after: page; }}
a {{ color: #0645ad; overflow-wrap: anywhere; }}
code {{ font-family: monospace; background: #f3f3f3; }}
ul {{ margin-top: .2em; }}
</style>
</head>
<body>
{cuerpo_html}
</body>
</html>"""
    SALIDA_HTML.write_text(pagina, encoding="utf-8")


def main():
    texto = MARKDOWN.read_text(encoding="utf-8")
    convertir_docx(texto)
    convertir_html(texto)
    print(f"DOCX: {SALIDA_DOCX}")
    print(f"HTML: {SALIDA_HTML}")
    print(f"DOCX bytes: {SALIDA_DOCX.stat().st_size}")


if __name__ == "__main__":
    main()
